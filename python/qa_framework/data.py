"""Portable JSON, CSV and Excel test data helpers."""
import csv
import json
from pathlib import Path


def read_json(path):
    return json.loads(Path(path).read_text(encoding="utf-8"))


def write_json(path, value):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, indent=2, ensure_ascii=False), encoding="utf-8")


def read_csv(path):
    with open(path, encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream))


def read_excel(path, sheet=None):
    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = (workbook[sheet] if sheet else workbook.active).iter_rows(values_only=True)
        headers = next(rows, None)
        return [dict(zip(headers, row)) for row in rows] if headers else []
    finally:
        workbook.close()


def write_excel(path, rows):
    from openpyxl import Workbook
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    try:
        rows = list(rows)
        if rows:
            headers = list(rows[0])
            workbook.active.append(headers)
            for row in rows:
                workbook.active.append([row.get(key) for key in headers])
        workbook.save(path)
    finally:
        workbook.close()
